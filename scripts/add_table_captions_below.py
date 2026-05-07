from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


CAPTIONS = [
    "Table 1. ISIC 2019 diagnostic label mapping used in this project.",
    "Table 2. Filtered ISIC 2019 class distribution after excluding UNK.",
    "Table 3. Train/validation/test split sizes and proxy robustness split sizes.",
    "Table 4. Training-split class counts and weighted cross-entropy class weights.",
    "Table 5. Model architecture summary for the baseline ViT and proposed Swin-Tiny.",
    "Table 6. Training configuration and reproducibility settings.",
    "Table 7. Performance impact of the three Swin improvement techniques relative to Swin without techniques.",
    "Table 8. Main test-set comparison across baseline, Swin variants, and foundation model.",
    "Table 9. Swin robustness proxy comparison on shared full-test / no-ruler / with-ruler splits.",
    "Table 10. Foundation-model comparison against the baseline ViT and best trained Swin model.",
    "Table 11. Summary of model-specific failure patterns and supporting evidence.",
    "Table 12. Individual team-member contributions.",
]


def make_caption_paragraph(anchor_table, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor_table._tbl.addnext(new_p)
    para = Paragraph(new_p, anchor_table._parent)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run(text)
    return para


def remove_existing_table_captions(doc: Document) -> None:
    to_remove = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("Table ") and ". " in text:
            to_remove.append(para._element)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def main() -> None:
    report_path = Path(r"d:\MS\Spring 2026\Deep Learning\Project 2\ruler-bias-vit\Final_Report.docx")
    doc = Document(str(report_path))

    if len(doc.tables) != len(CAPTIONS):
        raise RuntimeError(f"Expected {len(CAPTIONS)} tables, found {len(doc.tables)}")

    remove_existing_table_captions(doc)

    for table, caption in reversed(list(zip(doc.tables, CAPTIONS))):
        make_caption_paragraph(table, caption)

    doc.save(str(report_path))


if __name__ == "__main__":
    main()
