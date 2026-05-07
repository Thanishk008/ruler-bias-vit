from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "Final_Report.docx"


def find_paragraph(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"Paragraph not found: {startswith!r}")


def insert_paragraph_after(paragraph, text: str):
    new_p = paragraph.insert_paragraph_before("")
    paragraph._p.addnext(new_p._p)
    new_p.text = text
    return new_p


def set_cell(table, row: int, col: int, text: str) -> None:
    table.rows[row].cells[col].text = text


def main() -> None:
    doc = Document(str(DOCX_PATH))

    dataset_para = find_paragraph(doc, "The dataset used in this project is the public ISIC 2019 Kaggle dataset")
    dataset_para.text = (
        "The dataset used in this project is the public ISIC 2019 Kaggle dataset of dermoscopic skin lesion images "
        "with diagnostic ground-truth labels and metadata. After excluding UNK, the final dataset contains 25,331 JPG "
        "dermoscopic images with CSV label and metadata files. The source used in this project is the public Kaggle "
        "release at https://www.kaggle.com/datasets/andrewmvd/isic-2019. The data were prepared by removing UNK "
        "samples, encoding the eight retained classes, and creating stratified train/validation/test splits."
    )

    architecture_caption = find_paragraph(doc, "Architecture snapshot for the scratch-trained baseline ViT and the proposed Swin Transformer.")
    architecture_caption.text = "Figure 1. Architecture snapshot for the scratch-trained baseline ViT and the proposed Swin Transformer."

    figure_updates = {
        "Figure 1. Baseline confusion matrix.": "Figure 2. Baseline confusion matrix.",
        "Figure 2. Best-model confusion matrix (Technique 1).": "Figure 3. Best-model confusion matrix (Technique 1).",
        "Figure 3. Foundation-model confusion matrix.": "Figure 4. Foundation-model confusion matrix.",
        "Figure 4. Baseline precision-recall curves.": "Figure 5. Baseline precision-recall curves.",
        "Figure 5. Best-model precision-recall curves (Technique 1).": "Figure 6. Best-model precision-recall curves (Technique 1).",
        "Figure 6. Foundation-model precision-recall curves.": "Figure 7. Foundation-model precision-recall curves.",
        "Figure 7. Technique 1 training trend, epochs 1-10.": "Figure 8. Technique 1 training trend, epochs 1-10.",
        "Figure 8. Technique 1 training trend, epochs 11-50.": "Figure 9. Technique 1 training trend, epochs 11-50.",
        "Figure 9. Baseline Grad-CAM example: misclassified case, true BCC, predicted SCC.": "Figure 10. Baseline Grad-CAM example: misclassified case, true BCC, predicted SCC.",
        "Figure 10. Technique 1 Grad-CAM example: correctly classified case, true BCC, predicted BCC.": "Figure 11. Technique 1 Grad-CAM example: correctly classified case, true BCC, predicted BCC.",
        "Figure 11. Technique 1 Grad-CAM example: misclassified case, true BKL, predicted BCC.": "Figure 12. Technique 1 Grad-CAM example: misclassified case, true BKL, predicted BCC.",
    }
    for old, new in figure_updates.items():
        find_paragraph(doc, old).text = new

    foundation_para = find_paragraph(doc, "Likely reasons for CLIP’s weaker performance include")
    foundation_para.text = (
        "Likely reasons for CLIP’s weaker performance include lack of ISIC-specific supervision, coarse prompts for "
        "subtle dermatological categories, visually similar classes requiring domain-specific supervision, dataset "
        "imbalance, and domain shift between CLIP’s pretraining data and dermoscopic imaging. Quantitatively, the "
        "foundation model reaches macro precision 0.1683, macro recall 0.1430, and macro F1 0.1297, which is far "
        "below both the baseline ViT and the best trained Swin model. CLIP is a useful foundation comparison but not "
        "competitive with task-trained transformer models here."
    )

    gradcam_intro = find_paragraph(doc, "The following Grad-CAM examples provide qualitative evidence")
    gradcam_intro.text = (
        "The following Grad-CAM examples provide qualitative evidence about where the baseline and the best trained "
        "model focus for both correctly classified and misclassified cases. These examples support the error-analysis "
        "discussion by showing whether attention concentrates on clinically relevant lesion regions or drifts toward "
        "border context and other non-lesion structure."
    )

    gradcam_caption_1 = find_paragraph(doc, "Figure 10. Baseline Grad-CAM example: misclassified case, true BCC, predicted SCC.")
    gradcam_analysis_1 = insert_paragraph_after(
        gradcam_caption_1,
        "In Figure 10, the baseline misclassified BCC as SCC, illustrating that the baseline can fail even when the "
        "lesion belongs to a clinically important malignant class. This example is useful because it shows a failure "
        "case in which the baseline does not cleanly separate two visually similar diagnoses.",
    )
    gradcam_caption_2 = find_paragraph(doc, "Figure 11. Technique 1 Grad-CAM example: correctly classified case, true BCC, predicted BCC.")
    gradcam_analysis_2 = insert_paragraph_after(
        gradcam_caption_2,
        "In Figure 11, Technique 1 correctly classifies a BCC example, supporting the quantitative result that this "
        "configuration gives the best overall macro recall and macro F1. The example is consistent with the intended "
        "effect of border-focused preprocessing and synthetic ruler augmentation: the model is encouraged to rely more "
        "on lesion-relevant image content than on border shortcuts.",
    )
    gradcam_caption_3 = find_paragraph(doc, "Figure 12. Technique 1 Grad-CAM example: misclassified case, true BKL, predicted BCC.")
    gradcam_analysis_3 = insert_paragraph_after(
        gradcam_caption_3,
        "In Figure 12, Technique 1 still fails on a BKL-to-BCC confusion, showing that the strongest model is not "
        "error-free and that clinically subtle classes remain challenging. This qualitative example matches the class-"
        "wise metrics, where minority and visually similar classes still show limited precision despite the overall "
        "improvement in macro metrics.",
    )

    team_para = find_paragraph(doc, "The team used a shared problem definition, common dataset preparation")
    team_para.text = (
        "The team used a shared problem definition, common dataset preparation, and consistent train/validation/test "
        "splits so that all transformer approaches were evaluated under the same protocol. All members used the same "
        "shared ISIC 2019 split, baseline model, and evaluation setup for fair comparison across architectures. "
        "Coordination was handled through weekly Microsoft Teams meetings, shared GitHub repositories for version "
        "control, and shared report artifacts so that experiments, metrics, plots, and final written results could be "
        "integrated consistently. Collectively, the team established the common dataset pipeline, the shared baseline "
        "and evaluation framework, the cross-model comparison structure, and the final report, presentation, and video "
        "deliverables."
    )

    contributions_table = doc.tables[11]
    set_cell(
        contributions_table,
        1,
        1,
        "Implemented and evaluated the Segmented ViT approach, analyzed its three improvement-technique variants, "
        "integrated the Segmented ViT results into the cross-model comparison, and contributed to the final report, "
        "presentation, and video deliverables.",
    )
    set_cell(
        contributions_table,
        2,
        1,
        "Implemented and evaluated the Swin Transformer approach, including the three ruler/border-mitigation "
        "techniques used in this report, prepared the strongest trained-model analysis, and contributed to the shared "
        "experimental integration and final project deliverables.",
    )
    set_cell(
        contributions_table,
        3,
        1,
        "Implemented and evaluated the DeiT approach, analyzed its performance relative to the shared baseline, "
        "contributed DeiT results to the final comparison, and contributed to the final report, presentation, and "
        "video deliverables.",
    )

    robustness_para = find_paragraph(doc, "Because the robustness_comparison.csv files contain identical Full Test, No Ruler, and With Ruler values")
    robustness_para.text = (
        "Because the robustness_comparison.csv files contain identical Full Test, No Ruler, and With Ruler values for "
        "each Swin run, the table below reports the shared metric value once per run rather than repeating three "
        "identical columns. These results are interpreted only as consistency checks and not as a true ruler/no-ruler "
        "robustness experiment."
    )

    doc.save(str(DOCX_PATH))


if __name__ == "__main__":
    main()
