from __future__ import annotations

import math
from pathlib import Path
import sys

import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt
from matplotlib.patches import FancyBboxPatch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.models.baseline_vit import BaselineViT
from src.models.swin_transformer import SwinTransformer


DOCX_PATH = ROOT / "Final_Report.docx"
ASSET_DIR = ROOT / "outputs" / "report_assets"
ARCH_IMAGE = ASSET_DIR / "model_architecture_summary.png"
ARCH_TEXT = ASSET_DIR / "model_architecture_summary.txt"


def count_params(model) -> tuple[int, int]:
    total = sum(p.numel() for p in model.parameters())
    trainable = sum(p.numel() for p in model.parameters() if p.requires_grad)
    return total, trainable


def fmt_millions(value: int) -> str:
    return f"{value / 1_000_000:.2f}M"


def generate_architecture_assets() -> dict[str, dict[str, str]]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    baseline_wrapper = BaselineViT(num_classes=8)
    swin_wrapper = SwinTransformer(num_classes=8)
    baseline = baseline_wrapper.model
    swin = swin_wrapper.model

    baseline_total, baseline_trainable = count_params(baseline_wrapper)
    swin_total, swin_trainable = count_params(swin_wrapper)

    baseline_summary = {
        "model": "Baseline ViT-B/16",
        "patching": "16 x 16 patches -> 14 x 14 = 196 image tokens + 1 CLS token",
        "depth": f"{len(baseline.blocks)} transformer blocks",
        "embed": f"embed dim {baseline.embed_dim}, {baseline.blocks[0].attn.num_heads} attention heads, MLP hidden dim {baseline.blocks[0].mlp.fc1.out_features}",
        "params": fmt_millions(baseline_total),
        "trainable": fmt_millions(baseline_trainable),
        "output": "8-class linear head -> 8 logits",
    }

    stage_depths = [len(stage.blocks) for stage in swin.layers]
    stage_dims = [stage.blocks[0].norm1.normalized_shape[0] for stage in swin.layers]
    stage_heads = [stage.blocks[0].attn.num_heads for stage in swin.layers]
    stage_desc = ", ".join(
        f"S{i + 1}: depth {d}, dim {dim}, heads {heads}"
        for i, (d, dim, heads) in enumerate(zip(stage_depths, stage_dims, stage_heads))
    )
    swin_summary = {
        "model": "Proposed Swin-Tiny",
        "patching": "4 x 4 patch embedding with hierarchical patch merging",
        "depth": "4 stages with depths [2, 2, 6, 2]",
        "embed": f"stage dims [96, 192, 384, 768], heads [3, 6, 12, 24], window size {swin.layers[0].blocks[0].attn.window_size[0]}",
        "params": fmt_millions(swin_total),
        "trainable": fmt_millions(swin_trainable),
        "output": "global average pooling + 8-class head -> 8 logits",
        "stages": stage_desc,
    }

    with ARCH_TEXT.open("w", encoding="utf-8") as f:
        f.write("Model-summary style architecture snapshot\n\n")
        for summary in (baseline_summary, swin_summary):
            f.write(f"{summary['model']}\n")
            for key in ("patching", "depth", "embed", "params", "trainable", "output"):
                f.write(f"- {key}: {summary[key]}\n")
            if "stages" in summary:
                f.write(f"- stages: {summary['stages']}\n")
            f.write("\n")

    fig = plt.figure(figsize=(14, 8.5), facecolor="white")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.text(
        0.5,
        0.96,
        "Model-Summary Style Architecture Snapshot",
        ha="center",
        va="center",
        fontsize=18,
        fontweight="bold",
    )
    ax.text(
        0.5,
        0.92,
        "Generated from the actual PyTorch/timm model definitions used in this project",
        ha="center",
        va="center",
        fontsize=11,
        color="#444444",
    )

    def add_box(x: float, y: float, w: float, h: float, title: str, lines: list[str], fc: str) -> None:
        patch = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.012,rounding_size=0.018",
            linewidth=1.5,
            edgecolor="#1f2937",
            facecolor=fc,
        )
        ax.add_patch(patch)
        ax.text(x + 0.02, y + h - 0.04, title, ha="left", va="top", fontsize=12, fontweight="bold")
        text_y = y + h - 0.09
        for line in lines:
            ax.text(x + 0.02, text_y, line, ha="left", va="top", fontsize=10)
            text_y -= 0.052

    add_box(
        0.04,
        0.12,
        0.40,
        0.72,
        "Baseline ViT-B/16",
        [
            "Input: RGB image, 224 x 224",
            "Patch embedding: Conv2d kernel/stride 16 -> 196 patches",
            "Token sequence: 196 image tokens + 1 CLS token",
            "Encoder depth: 12 transformer blocks",
            "Attention: 12 heads, embed dim 768",
            "MLP hidden dim: 3072",
            f"Parameters: {baseline_summary['params']} total / {baseline_summary['trainable']} trainable",
            "Classifier: linear head -> 8 logits",
            "Purpose in report: shared scratch-trained baseline",
        ],
        "#e7f0fb",
    )

    add_box(
        0.56,
        0.12,
        0.40,
        0.72,
        "Proposed Swin-Tiny",
        [
            "Input: RGB image, 224 x 224",
            "Patch embedding: Conv2d kernel/stride 4",
            "Hierarchy: 4 stages with patch merging between stages",
            "Stage 1: depth 2, dim 96, 3 heads",
            "Stage 2: depth 2, dim 192, 6 heads",
            "Stage 3: depth 6, dim 384, 12 heads",
            "Stage 4: depth 2, dim 768, 24 heads",
            "Window attention: 7 x 7 shifted windows",
            f"Parameters: {swin_summary['params']} total / {swin_summary['trainable']} trainable",
            "Classifier: global average pooling + 8-class head",
        ],
        "#e8f6ee",
    )

    ax.annotate("", xy=(0.56, 0.50), xytext=(0.44, 0.50), arrowprops=dict(arrowstyle="<->", lw=1.5, color="#334155"))
    ax.text(0.5, 0.53, "Architectural contrast", ha="center", va="bottom", fontsize=9, color="#334155")
    ax.text(
        0.5,
        0.485,
        "flat global-token ViT",
        ha="right",
        va="center",
        fontsize=9,
        color="#334155",
    )
    ax.text(
        0.5,
        0.485,
        " vs hierarchical shifted-window Swin",
        ha="left",
        va="center",
        fontsize=9,
        color="#334155",
    )

    fig.savefig(ARCH_IMAGE, dpi=220, bbox_inches="tight")
    plt.close(fig)

    return {"baseline": baseline_summary, "swin": swin_summary}


def find_paragraph_by_prefix(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Could not find paragraph starting with: {prefix!r}")


def has_paragraph_prefix(doc: Document, prefix: str) -> bool:
    return any(paragraph.text.strip().startswith(prefix) for paragraph in doc.paragraphs)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


def insert_table_after(paragraph: Paragraph, rows: int, cols: int, style: str = "Table Grid") -> Table:
    doc = paragraph._parent.part.document
    table = doc.add_table(rows=rows, cols=cols)
    table.style = style
    paragraph._p.addnext(table._tbl)
    return table


def insert_paragraph_after_table(table: Table, text: str = "", style: str | None = None) -> Paragraph:
    doc = table._parent.part.document
    paragraph = doc.add_paragraph()
    if text:
        paragraph.add_run(text)
    if style:
        paragraph.style = style
    table._tbl.addnext(paragraph._p)
    return paragraph


def add_centered_picture_after(paragraph: Paragraph, image_path: Path, width_inches: float = 6.8) -> Paragraph:
    pic_para = insert_paragraph_after(paragraph)
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_para.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    return pic_para


def find_table_by_header(doc: Document, first_cell_text: str) -> Table:
    for table in doc.tables:
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip() == first_cell_text:
            return table
    raise RuntimeError(f"Could not find table with first header cell: {first_cell_text!r}")


def fill_table(table: Table, rows: list[list[str]]) -> None:
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = value


def main() -> None:
    summaries = generate_architecture_assets()
    doc = Document(str(DOCX_PATH))

    dataset_para = find_paragraph_by_prefix(doc, "The dataset used in this project is the public ISIC 2019 Kaggle dataset")
    dataset_para.text = (
        "The dataset used in this project is the public ISIC 2019 Kaggle dataset of dermoscopic skin lesion images with "
        "diagnostic ground-truth labels and metadata. After excluding UNK, the final dataset contains 25,331 JPG "
        "dermoscopic images with CSV label and metadata files. The data were prepared by removing UNK samples, encoding "
        "the eight retained classes, and creating stratified train/validation/test splits."
    )

    if not has_paragraph_prefix(doc, "Full filtered class distribution:"):
        dist_intro = insert_paragraph_after(
            dataset_para,
            "Full filtered class distribution:",
        )
        dist_table = insert_table_after(dist_intro, rows=9, cols=3)
        fill_table(
            dist_table,
            [
                ["Label", "Class", "Count in filtered ISIC 2019 dataset"],
                ["0", "MEL", "4522"],
                ["1", "NV", "12875"],
                ["2", "BCC", "3323"],
                ["3", "AK", "867"],
                ["4", "BKL", "2624"],
                ["5", "DF", "239"],
                ["6", "VASC", "253"],
                ["7", "SCC", "628"],
            ],
        )

    split_heading = find_paragraph_by_prefix(doc, "3.2 ")
    split_heading.text = "3.2 Training split imbalance and class weights"
    split_para = find_paragraph_by_prefix(doc, "The training split is strongly imbalanced:")
    split_para.text = (
        "The full filtered dataset is highly imbalanced, and the training split preserves that imbalance after "
        "stratification. NV dominates the dataset, while DF, VASC, SCC, and AK are much smaller. Because class weights "
        "are computed from the training split, the table below reports the training counts and resulting class weights "
        "used for weighted cross-entropy."
    )

    train_transform_para = find_paragraph_by_prefix(doc, "The trained ViT and Swin models use ImageNet-style input normalization")
    train_transform_para.text = (
        "For the trained ViT and Swin models, the training transform applies RandomHorizontalFlip(p=0.5), "
        "RandomVerticalFlip(p=0.5), RandomRotation(20), ColorJitter(0.2, 0.2, 0.2, 0.1), Resize((224, 224)), "
        "ToTensor(), and ImageNet normalization with mean [0.485, 0.456, 0.406] and std [0.229, 0.224, 0.225]. "
        "Validation and test use deterministic preprocessing: Resize(224), CenterCrop(224), ToTensor(), and the same "
        "ImageNet normalization. Class weights are computed from the training split and passed to cross-entropy loss to "
        "reduce the effect of class imbalance."
    )

    foundation_transform_para = find_paragraph_by_prefix(doc, "The foundation model uses CLIP-specific preprocessing and normalization")
    foundation_transform_para.text = (
        "The foundation model uses CLIP-specific preprocessing: Resize(224), CenterCrop(224), ToTensor(), and CLIP "
        "normalization with mean [0.48145466, 0.4578275, 0.40821073] and std [0.26862954, 0.26130258, 0.27577711]. "
        "Data leakage is avoided by creating the train, validation, and test splits before model training and by "
        "selecting the best checkpoint using validation macro recall, not test performance. The test set is used only "
        "for final evaluation. The duplicated no-ruler and with-ruler proxy splits are derived from the same test set "
        "and are not used for training."
    )

    baseline_perf_para = find_paragraph_by_prefix(doc, "This satisfies the project requirement that proposed models be trained from scratch.")
    baseline_perf_para.text = (
        "This satisfies the project requirement that proposed models be trained from scratch. The baseline is trained "
        "on the ISIC 2019 training split with an 8-class classification head. Its final test performance is macro "
        "precision 0.3849, macro recall 0.4867, and macro F1 0.4048. It provides a transformer-based reference point "
        "for comparison against the proposed Swin Transformer and improvement techniques."
    )

    proposed_arch_para = find_paragraph_by_prefix(doc, "The proposed model is a Swin Transformer using")
    if not has_paragraph_prefix(doc, "Architecture Summary"):
        arch_heading = insert_paragraph_after(proposed_arch_para, "Architecture Summary")
        arch_heading.runs[0].bold = True
        arch_text = insert_paragraph_after(
            arch_heading,
            "The summaries below were generated from the actual PyTorch/timm model definitions used in this project and "
            "provide a report-friendly alternative to a TensorFlow-style model.summary() view.",
        )
        arch_table = insert_table_after(arch_text, rows=3, cols=6)
        fill_table(
            arch_table,
            [
                ["Model", "Input / patching", "Core depth", "Attention / dimensions", "Parameters", "Output"],
                [
                    summaries["baseline"]["model"],
                    "224 x 224 RGB; 16 x 16 patch embedding -> 196 image tokens + CLS",
                    "12 transformer blocks",
                    "embed dim 768; 12 heads; MLP 3072",
                    summaries["baseline"]["params"],
                    "8 logits",
                ],
                [
                    summaries["swin"]["model"],
                    "224 x 224 RGB; 4 x 4 patch embedding with patch merging",
                    "4 stages with depths [2, 2, 6, 2]",
                    "dims [96, 192, 384, 768]; heads [3, 6, 12, 24]; window size 7",
                    summaries["swin"]["params"],
                    "8 logits",
                ],
            ],
        )
        pic_para = insert_paragraph_after_table(arch_table)
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_para.add_run().add_picture(str(ARCH_IMAGE), width=Inches(6.8))
        caption = insert_paragraph_after(
            pic_para,
            "Architecture snapshot for the scratch-trained baseline ViT and the proposed Swin Transformer.",
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    main_results_heading = find_paragraph_by_prefix(doc, "7.1 Main Test Results")
    if not has_paragraph_prefix(doc, "The table below reports the primary model-selection metrics"):
        insert_paragraph_after(
            main_results_heading,
            "The table below reports the primary model-selection metrics: macro precision, macro recall, and macro F1. "
            "Accuracy is discussed separately in the robustness and failure-analysis sections because this report "
            "emphasizes class-balanced evaluation under strong class imbalance.",
        )

    robustness_para = find_paragraph_by_prefix(doc, "The robustness comparison shows identical full-test, no-ruler, and with-ruler metrics")
    robustness_para.text = (
        "Because the robustness_comparison.csv files contain identical Full Test, No Ruler, and With Ruler values for "
        "each Swin run, the table below reports the shared value once per run. These results are interpreted only as "
        "consistency checks and not as a true ruler/no-ruler robustness experiment."
    )

    team_para = find_paragraph_by_prefix(doc, "The team used a shared problem definition, common dataset preparation")
    team_para.text = (
        "The team used a shared problem definition, common dataset preparation, and consistent train/validation/test "
        "splits so that all transformer approaches were evaluated under the same protocol. All members used the same "
        "shared ISIC 2019 split, baseline model, and evaluation setup for fair comparison across architectures. "
        "Coordination was handled through weekly Microsoft Teams meetings, shared GitHub repositories for version "
        "control, and shared report artifacts so that experiments, metrics, plots, and final written results could be "
        "integrated consistently. Collectively, the team produced the shared baseline/foundation comparison setup, the "
        "final report, the final presentation, and the presentation video."
    )

    contributions_table = find_table_by_header(doc, "Team Member")
    contributions_rows = [
        ["Team Member", "Contributions"],
        [
            "Jinali",
            "Implemented and evaluated the Segmented ViT approach, analyzed its three improvement-technique variants, "
            "and integrated the Segmented ViT results into the cross-model comparison.",
        ],
        [
            "Thanishk",
            "Implemented and evaluated the Swin Transformer approach, including the three ruler/border-mitigation "
            "techniques used in this report, and prepared the strongest trained-model analysis.",
        ],
        [
            "Tommy",
            "Implemented and evaluated the DeiT approach, analyzed its performance relative to the shared baseline, "
            "and contributed DeiT results to the final comparison.",
        ],
    ]
    fill_table(contributions_table, contributions_rows)

    conclusion_para = find_paragraph_by_prefix(doc, "A key lesson learned is that bias-mitigation methods must be evaluated")
    conclusion_para.text = (
        "A key lesson learned is that bias-mitigation methods must be evaluated with both class-imbalance-aware metrics "
        "and dataset assumptions in mind: a technique can improve recall while still hurting overall F1, and a bias "
        "claim is weaker when explicit ruler annotations are unavailable. The main limitation is that the processed "
        "ISIC 2019 setup does not include explicit ruler/no-ruler annotations. Therefore, this project is best "
        "interpreted as a study of ruler-like border artifact mitigation rather than a definitive ruler-bias "
        "measurement study. Future work should use a curated ruler-annotated test set, tune the strength of border "
        "masking and attention regularization, and explore more clinically grounded prompts or fine-tuned medical "
        "foundation models. Despite this limitation, the project demonstrates a complete transformer classification "
        "pipeline with scratch-trained models, consistent splits, multiple improvement techniques, foundation-model "
        "comparison, quantitative metrics, and visual analysis."
    )

    doc.save(str(DOCX_PATH))


if __name__ == "__main__":
    main()
