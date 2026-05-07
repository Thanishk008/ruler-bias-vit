from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SRC_DOCX = ROOT / "Final_Report.docx"
OUT_DOCX = ROOT / "Final_Report_clean.docx"

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def get_para_text(paragraph_el) -> str:
    return "".join(t.text or "" for t in paragraph_el.findall(".//w:t", NS)).strip()


def get_images_for_paragraph(paragraph_el, rels_map: dict[str, str]) -> list[str]:
    images = []
    for blip in paragraph_el.findall(".//a:blip", NS):
        rel_id = blip.get(f"{{{NS['r']}}}embed")
        if rel_id and rel_id in rels_map:
            images.append(rels_map[rel_id])
    return images


def style_for_text(text: str) -> str | None:
    if not text:
        return None
    if text == "Fixing Ruler-Like Border Artifact Bias in Vision Transformers for Skin Lesion Classification":
        return "Title"
    if re.match(r"^\d+\.\s", text):
        return "Heading 1"
    if re.match(r"^\d+\.\d+\s", text):
        return "Heading 2"
    if text in {
        "Architecture Summary",
        "Technique Impact Summary",
        "Confusion Matrices",
        "Precision-Recall Curves",
        "Epoch Trends",
        "Grad-CAM Examples",
    }:
        return "Heading 2"
    return None


def add_paragraph(doc: Document, text: str) -> None:
    style = style_for_text(text)
    para = doc.add_paragraph(style=style)
    para.add_run(text)
    if text.startswith("Figure "):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(doc: Document, table_el) -> None:
    rows = table_el.findall("./w:tr", NS)
    if not rows:
        return
    col_count = max(len(row.findall("./w:tc", NS)) for row in rows)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    for row_el in rows:
        row = table.add_row().cells
        cells = row_el.findall("./w:tc", NS)
        for idx in range(col_count):
            row[idx].text = ""
        for idx, cell_el in enumerate(cells):
            cell_text = " ".join(
                "".join(t.text or "" for t in p.findall(".//w:t", NS)).strip()
                for p in cell_el.findall(".//w:p", NS)
                if "".join(t.text or "" for t in p.findall(".//w:t", NS)).strip()
            ).strip()
            row[idx].text = cell_text


def main() -> None:
    with zipfile.ZipFile(SRC_DOCX) as zf:
        document_xml = etree.fromstring(zf.read("word/document.xml"))
        rels_xml = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
        rels_map = {}
        for rel in rels_xml:
            rel_id = rel.get("Id")
            target = rel.get("Target")
            rel_type = rel.get("Type", "")
            if rel_id and "image" in rel_type and target:
                rels_map[rel_id] = f"word/{target}"

        body = document_xml.find(".//w:body", NS)
        if body is None:
            raise RuntimeError("Could not find document body")

        doc = Document()

        for child in body:
            local_name = etree.QName(child).localname
            if local_name == "p":
                text = get_para_text(child)
                images = get_images_for_paragraph(child, rels_map)

                if text:
                    add_paragraph(doc, text)
                elif images:
                    doc.add_paragraph("")

                for image_path in images:
                    image_bytes = zf.read(image_path)
                    para = doc.add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run()
                    run.add_picture(BytesIO(image_bytes), width=Inches(6.5))
            elif local_name == "tbl":
                add_table(doc, child)

        doc.save(OUT_DOCX)
        print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
