from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


CAPTIONS = [
    "Table 1. ISIC 2019 diagnostic label mapping used in this project.",
    "Table 2. Filtered ISIC 2019 class distribution after excluding UNK.",
    "Table 3. Train/validation/test split sizes and proxy robustness split sizes.",
    "Table 4. Training-split class counts and weighted cross-entropy class weights.",
    "Table 5. Model architecture summary for the baseline ViT and proposed Swin-Tiny.",
    "Table 6. Training configuration and reproducibility settings.",
    "Table 7. Performance impact of the three Swin improvement techniques relative to Swin without techniques.",
    "Table 8. Main test-set comparison across baseline, Swin variants, and foundation model.",
    "Table 9. Swin robustness proxy comparison on shared full-test / no-ruler / with-ruler splits.",
    "Table 10. Foundation-model comparison against the baseline ViT and best trained Swin model.",
    "Table 11. Summary of model-specific failure patterns and supporting evidence.",
    "Table 12. Individual team-member contributions.",
]


def insert_paragraph_before_table(table, text: str, style_name: str = "Normal") -> Paragraph:
    tbl = table._tbl
    new_p = OxmlElement("w:p")
    tbl.addprevious(new_p)
    para = Paragraph(new_p, table._parent)
    para.style = style_name
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    ref_run = None
    for paragraph in table._parent.paragraphs:
        if paragraph.text.strip().startswith("Figure 1."):
            ref_run = paragraph.runs[0] if paragraph.runs else None
            break
    if ref_run is not None:
        run.bold = ref_run.bold
        run.italic = ref_run.italic
        if ref_run.font.name:
            run.font.name = ref_run.font.name
        if ref_run.font.size:
            run.font.size = ref_run.font.size
    return para


def main() -> None:
    report_path = Path(r"d:\MS\Spring 2026\Deep Learning\Project 2\ruler-bias-vit\Final_Report.docx")
    output_path = Path(r"d:\MS\Spring 2026\Deep Learning\Project 2\ruler-bias-vit\Final_Report_with_table_titles.docx")
    doc = Document(str(report_path))

    if len(doc.tables) != len(CAPTIONS):
        raise RuntimeError(f"Expected {len(CAPTIONS)} tables, found {len(doc.tables)}")

    existing = {p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Table ")}
    for caption in reversed(CAPTIONS):
        if caption in existing:
            continue

    for table, caption in zip(doc.tables, CAPTIONS):
        prev = table._tbl.getprevious()
        already_captioned = False
        if prev is not None and prev.tag.endswith("}p"):
            prev_para = Paragraph(prev, table._parent)
            if prev_para.text.strip().startswith("Table "):
                prev_para.text = caption
                prev_para.style = "Normal"
                prev_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                already_captioned = True
        if not already_captioned:
            insert_paragraph_before_table(table, caption)

    doc.save(str(output_path))


if __name__ == "__main__":
    main()
